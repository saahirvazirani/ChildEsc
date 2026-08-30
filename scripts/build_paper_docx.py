"""Build the full ChildEsc collaboration manuscript as a DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "paper" / "ChildEsc_Workshop_Paper.docx"


def set_font(run, name: str = "Times New Roman", size: float = 10, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=5.5, line=11):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(line)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "insideH", "bottom"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6" if edge != "insideH" else "2")
        node.set(qn("w:color"), "333333" if edge != "insideH" else "B7B7B7")
        borders.append(node)
    for edge in ("left", "right", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def set_table_widths(table, widths: list[int]):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def mark_table_header(row):
    row_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_pr.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))
    set_font(run, size=9)


def add_body(
    doc,
    text: str,
    bold_lead: str | None = None,
    *,
    size: float = 10,
    after: float = 5.5,
    line: float = 11,
):
    paragraph = doc.add_paragraph()
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_font(lead, size=size, bold=True)
    run = paragraph.add_run(text)
    set_font(run, size=size)
    set_paragraph_spacing(paragraph, after=after, line=line)
    return paragraph


def add_heading(doc, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_font(run, size=12 if level == 1 else 10, bold=True)
    paragraph.paragraph_format.keep_with_next = True
    set_paragraph_spacing(paragraph, before=8 if level == 1 else 5.5, after=3, line=12)
    return paragraph


def add_reference(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    run = paragraph.add_run(text)
    set_font(run, size=8.5)
    set_paragraph_spacing(paragraph, after=2, line=9)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.5)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = Pt(11)

    for name, size in (("Heading 1", 12), ("Heading 2", 10)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.line_spacing = Pt(20)
    title_run = title.add_run(
        "Who Gets Involved? ChildEsc Audits Escalation and Handoff Routing in Child-Facing AI"
    )
    set_font(title_run, size=17, bold=True)
    p_pr = title._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge, size in (("top", "32"), ("bottom", "8")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "8")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    p_pr.append(borders)

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(author, after=11)
    set_font(author.add_run("Anonymous Author(s)"), bold=True)

    abstract_heading = doc.add_paragraph()
    abstract_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(abstract_heading, before=11, after=5.5, line=12)
    set_font(abstract_heading.add_run("Abstract"), size=12, bold=True)
    abstract = doc.add_paragraph()
    abstract.paragraph_format.left_indent = Inches(0.5)
    abstract.paragraph_format.right_indent = Inches(0.5)
    set_paragraph_spacing(abstract, after=11)
    abstract_text = (
        "Child-facing AI increasingly mediates emotional support and may route young users toward real-world help. Yet action-only evaluations "
        "can reward correct urgency while obscuring whether a recommended recipient satisfies context-specific constraints. We introduce ChildEsc, a works-in-progress testbed "
        "of 80 LLM-assisted synthetic conversations in 20 counterfactual families, four ordered routing actions, and explicit permitted and "
        "forbidden targets. A provider-neutral harness audits action-plus-recipient decisions using the action-route gap between action accuracy "
        "and exact routing. Across eight complete, internally cache-replayed trials of three prompted systems, mean action accuracy ranges from 80.0% to 83.1% "
        "while exact routing ranges from 33.8% to 51.7%, yielding 29.2- to 49.4-point gaps; 42.7% to 71.2% of action-correct handoffs "
        "violate target constraints. Recipient constraints are therefore a distinct, inspectable safeguard dimension, but these results remain "
        "policy-exposed specification-conformance diagnostics: practitioner and youth validation have not occurred, and no result supports deployment."
    )
    set_font(abstract.add_run(abstract_text))

    add_heading(doc, "1 Introduction")
    add_body(doc, "Children use conversational AI for information, companionship, and emotional support, where a superficially safe response can still worsen a vulnerable situation (UNICEF, 2026; Cha et al., 2026). Binary respond/refuse policies are especially ill-suited to these interactions: refusal can close a rare help-seeking window, while unbounded engagement can reinforce dependency or miss imminent danger. Interviews with 19 youth-facing practitioners found that useful systems should gather context and bridge youth to tailored human support, but should not assume a parent is safe (Cha et al., 2026). This choice is already operational: OpenAI reports automatically placing users estimated to be under 18 or self-identifying as 13-17 into a teen experience; linked parents can receive safety notifications in limited high-risk situations (OpenAI, 2026). The deployment decision is therefore not simply \"is this content allowed?\" but what support action is proportionate now, and who is safe to involve?")
    add_body(doc, "Related benchmarks cover important parts of this problem. MinorBench and Safe-Child-LLM emphasize unsafe-request handling; KIDBench evaluates developmental response quality and trusted-adult redirection; CAREBench includes crisis referral and human-support behavior; KORA evaluates simulated multi-turn child interactions with LLM-judged safety and mechanism grades; and MindEval measures multi-turn adult mental-health support. CARE-Bench recently operationalized patient-facing medical triage as a four-label sequential current-action task, including whether more information is needed. Thus neither referral, ordered actions, timing, nor multi-turn child simulation is our novelty. ChildEsc's distinct contribution is recipient-constrained routing: each child-specific, severity-matched trajectory couples an ordered action with explicit permitted and forbidden handoff sets rather than scoring generated response quality. This maps directly to the workshop's restricted-data evaluation and deployment-safeguard themes.")
    add_body(doc, "We contribute: (1) a set-valued recipient-constrained routing task, instantiated as 80 synthetic trajectories with four actions and item-specific permitted and forbidden targets; (2) the action-route gap and an atomic failure taxonomy that expose action-correct but target-invalid handoffs; and (3) a provider-neutral, internally replayed and hash-documented audit across three prompted systems, plus an executable checksum and pre-specified validation gates. We do not claim benchmark validity, child-centeredness, or system superiority before independent review.")

    add_heading(doc, "2 ChildEsc testbed")
    add_body(doc, "ChildEsc applies to text-based consumer assistants, companions, tutors, and wellness or mental-health-adjacent systems used by or accessible to people under 18. It evaluates an application-layer router before or alongside response generation. It does not diagnose, recommend treatment, or automatically contact any person or service.", "Scope. ")
    add_body(doc, "Under a written specification, an LLM coding assistant drafted 20 scenario families across eight vulnerability domains. Each family contains a shared context and four severity cues, expanded deterministically into three-user-turn conversations. No text derives from a real child, hotline transcript, clinical record, or explicit abuse material. A dated register maps domain-level rationales to public sources, but this retrospective mapping is source-informed rather than item-level grounding. The same workflow drafted labels and rules, making all results policy-exposed and unsuitable for estimating generalization.", "Construction and provenance. ")
    add_body(doc, "The ordinal action space is continue_support (0), boundary_redirect (1), human_support (2), and urgent_handoff (3). Actions encode intervention and routing obligation rather than complete response style: continue_support requires neither boundary redirection nor handoff; boundary_redirect requires boundary-setting without human handoff; human_support requires a non-urgent human pathway; and urgent_handoff prioritizes immediate connection. Supportive language may accompany any action. Escalated items permit one or more target classes. A target names a recommended recipient class and does not authorize autonomous contact or data disclosure. Forbidden targets use the same vocabulary a router can emit; safe_adult is forbidden when the conversation implicates an adult, while safe_adult_not_implicated may be permitted.", "Actions and targets. ")
    add_body(doc, "For reference and predicted actions g_i and p_i, and permitted, forbidden, and predicted target sets P_i, F_i, and T_i, target validity V_i requires either empty P_i and T_i, or nonempty T_i contained in P_i and disjoint from F_i. Exact routing is 1[p_i = g_i]V_i. Under-escalation is conditioned on the 58 items requiring human or urgent support. The action-route gap is action accuracy minus exact-route accuracy. The artifact also computes counterfactual sensitivity and monotonicity. Reaction delay is omitted because v0.1 fixes the decisive cue turn.", "Failure-oriented metrics. ")

    add_heading(doc, "3 Diagnostic implementation and audit")
    add_body(doc, "ChildEsc-Rules is a transparent policy-exposed checksum; two simpler baselines remain in the appendix. A provider-neutral harness sends only conversation text and requests the same action-target schema. We evaluate dated Claude Sonnet 5, Gemini 3.1 Pro Preview, and GPT-5.5 backends through OpenRouter. Each is pinned to its first-party endpoint with fallbacks disabled, low reasoning, a 1,024-token cap, omitted temperature, unique trial IDs, append-only attempts, and content-addressed caches. Layer 1 selects an action and recipient, and ChildEsc evaluates only that routing decision. Layer 2 generates supportive language, which requires a separate rubric and governed human validation.", "Policies and evaluator. ")

    caption = doc.add_paragraph()
    caption.paragraph_format.keep_with_next = True
    set_paragraph_spacing(caption, before=5.5, after=5.5)
    set_font(caption.add_run("Table 1: Prompted-system diagnostics on 80 synthetic items."), bold=True)
    headers = ["System", "Trials", "Action", "Exact", "Gap", "TgtFail", "RouteVar"]
    data = [
        ["Claude Sonnet 5", "3", ".808", ".517", ".292", ".427", ".225"],
        ["Gemini 3.1 Pro Preview", "3", ".800", ".392", ".408", ".612", ".513"],
        ["GPT-5.5", "2", ".831", ".338", ".494", ".712", ".263"],
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    mark_table_header(table.rows[0])
    widths = [Inches(1.22)] + [Inches(0.71)] * 6
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "E7E7E7")
        set_cell_margins(cell)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cell.paragraphs[0].add_run(header), size=8.5, bold=True)
    for row_index, values in enumerate(data):
        row = table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = widths[index]
            set_cell_margins(cell)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_font(cell.paragraphs[0].add_run(value), size=8.5)
    set_table_widths(table, [1760] + [1026] * 6)
    set_table_borders(table)
    note = doc.add_paragraph()
    set_paragraph_spacing(note, after=5.5, line=9)
    set_font(note.add_run("Values are trial means. TgtFail conditions target-constraint invalidity on action-correct handoffs; RouteVar is full-route instability."), size=8, italic=True)

    add_body(doc, "Across systems, mean action accuracy ranges from 80.0% to 83.1% while exact routing ranges from 33.8% to 51.7%, producing 29.2- to 49.4-point action-route gaps. Among action-correct handoffs, 42.7% to 71.2% violate target constraints. Action instability is 1.3% to 6.3% of items, whereas full-route instability is 22.5% to 51.3%. Urgent recall is 97.3% to 100%. The deterministic checksum independently shows an 11.3-point gap and 9/37 action-correct target failures.", "Action-route gap. ")
    add_body(doc, "Every one of the 247 action-correct target failures includes an extra unpermitted class; none is caused only by an omitted target, and only two also intersect a forbidden class. The discrepancy therefore primarily reflects over-broad provisional recipient sets and specification nonconformance, not demonstrated routing to substantively unsafe people. In one self-harm item, a system selects the correct human-support action but adds an unpermitted crisis-service target.", "Failure decomposition. ")
    decomposition_caption = doc.add_paragraph()
    decomposition_caption.paragraph_format.keep_with_next = True
    set_paragraph_spacing(decomposition_caption, before=4, after=4)
    set_font(decomposition_caption.add_run("Table 2: Failure decomposition across complete trials."), bold=True)
    decomposition_headers = ["System", "Handoffs", "Failures", "Unpermitted", "Missing", "Forbidden"]
    decomposition_data = [
        ["Claude Sonnet 5", "164", "70", "70", "0", "0"],
        ["Gemini 3.1 Pro Preview", "160", "98", "98", "0", "2"],
        ["GPT-5.5", "111", "79", "79", "0", "0"],
    ]
    decomposition_table = doc.add_table(rows=1, cols=len(decomposition_headers))
    decomposition_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    decomposition_table.autofit = False
    mark_table_header(decomposition_table.rows[0])
    decomposition_widths = [Inches(1.4), Inches(0.72), Inches(0.72), Inches(0.92), Inches(0.72), Inches(0.72)]
    for index, header in enumerate(decomposition_headers):
        cell = decomposition_table.rows[0].cells[index]
        cell.width = decomposition_widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "E7E7E7")
        set_cell_margins(cell, start=30, end=30)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cell.paragraphs[0].add_run(header), size=7.5, bold=True)
    for values in decomposition_data:
        row = decomposition_table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = decomposition_widths[index]
            set_cell_margins(cell, start=30, end=30)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_font(cell.paragraphs[0].add_run(value), size=7.5)
    set_table_widths(decomposition_table, [2016, 1037, 1037, 1325, 1037, 1037])
    set_table_borders(decomposition_table)
    decomposition_note = doc.add_paragraph()
    set_paragraph_spacing(decomposition_note, after=5.5, line=9)
    set_font(decomposition_note.add_run("Handoffs counts action-correct items requiring targets; failure reasons may co-occur. Missing reports failures caused only by emitting no target."), size=8, italic=True)
    add_body(doc, "These tests probe the frozen checksum without retuning it and are not held-out evidence. Latest-turn-only input under-escalates all 58 escalation-required items. Disabling only adult-implication inference leaves actions unchanged but produces a forbidden safe-adult target in 7/12 implicated-adult cases. Uppercasing preserves every route; doubling spaces reduces action agreement to 45.0% and urgent recall from 75.7% to 10.8%.", "Diagnostic stress tests. ")
    add_body(doc, "Version 1.1 was pre-specified after a failed pilot and produced three complete Claude trials. Version 1.2 changed no benchmark item, label, prompt, schema, request setting, score, or retry rule; before new calls it fixed Gemini/GPT identities, endpoint snapshots, spend and inclusion rules, and bounded claim language based on whether zero, one, or both comparison systems showed a positive gap. All complete results are disclosed. Three Gemini trials and two GPT trials completed; one GPT trial failed at 79/80 and remains unscored without retry. All scored runs were replayed internally from retained caches with matching hashes; the supplement distributes hashes and aggregates, not raw responses or caches.", "Protocol and replay. ")

    add_heading(doc, "4 Source-informed relational contract audit")
    add_body(doc, "Nine matched, non-procedural probe families were frozen after v0.1 to test how routes change with one operational factor. Public sources motivate each relation but do not determine labels. Because the developers had inspected the router, this is a post-freeze mechanistic audit, not held-out evaluation.")
    contract_caption = doc.add_paragraph()
    contract_caption.paragraph_format.keep_with_next = True
    set_paragraph_spacing(contract_caption, before=4, after=4)
    set_font(contract_caption.add_run("Table 3: Frozen contract audit."), bold=True)
    contract_headers = ["Controlled factor", "Base to variant", "Pass", "Family"]
    contract_data = [
        ["Temporal immediacy", "Urgent to Urgent", "3/5", "fail"],
        ["Resolved negation", "Urgent to Continue", "3/3", "pass"],
        ["First/third person", "Urgent to Human", "4/4", "pass"],
        ["Implicated adult", "Human to Human", "6/6", "pass"],
        ["AI dependency", "Continue to Human", "4/4", "pass"],
        ["Coercive secrecy", "Continue to Human", "4/4", "pass"],
        ["Online threat", "Boundary to Urgent", "5/5", "pass"],
        ["Medical deterioration", "Boundary to Urgent", "4/5", "fail"],
        ["Explicit/implicit age", "Human to Human", "3/3", "pass"],
    ]
    contract_table = doc.add_table(rows=1, cols=len(contract_headers))
    contract_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    contract_table.autofit = False
    mark_table_header(contract_table.rows[0])
    contract_widths = [Inches(1.75), Inches(1.75), Inches(0.65), Inches(0.65)]
    for index, header in enumerate(contract_headers):
        cell = contract_table.rows[0].cells[index]
        cell.width = contract_widths[index]
        set_cell_shading(cell, "E7E7E7")
        set_cell_margins(cell, top=20, bottom=20, start=40, end=40)
        set_font(cell.paragraphs[0].add_run(header), size=7.5, bold=True)
    for values in contract_data:
        row = contract_table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = contract_widths[index]
            set_cell_margins(cell, top=10, bottom=10, start=40, end=40)
            set_font(cell.paragraphs[0].add_run(value), size=7.5)
    set_table_borders(contract_table)
    add_body(doc, "Overall, 36/39 assertions and 7/9 families pass. The phrase 'safe right now' triggers urgency despite explicit negation because 'right now' is matched globally; physical deterioration triggers urgency but omits health_professional. Both failures are retained and the checksum is not modified. The supported use is auditing implementation conformance to an inspectable routing contract, not ranking models or certifying safety.")
    add_body(doc, "ChildEsc accepts any router that emits the two-field action-target schema and produces item-level failures, family-grouped intervals, repeatability diagnostics, and a hash-documented audit trail. This supports pre-release comparison of routing-policy versions, diagnosis of action-correct recipient errors, and review of implicated-adult defaults. It is an executable object for safeguard deliberation, not a deployment gate.", "Technical utility. ")
    add_body(doc, "An incomplete informal editorial read by a non-professional individual with mental-health knowledge revealed a dual interpretation of continue_support: supportive conversation may continue even when a handoff is required. Actions are therefore defined by intervention and routing obligations, not full response style. The read also foregrounded missing medical/immediacy context and jurisdiction-dependent recipients. It is not validation, changes no v0.1 label, and motivates an explicit abstain or clarify state and uncertainty-valued targets for v0.2.", "Review implications. ")

    add_heading(doc, "5 Validity, impact, and next evaluation")
    add_body(doc, "The central validity threat is circularity: scenarios, labels, lexical rules, probes, and prose were drafted in one LLM-assisted workflow. Source traceability improves auditability, not validity. The three prompted systems use the same provisional specification and prompt; this is construct replication, not an independent holdout or defensible leaderboard. The English convenience set omits dialect, code-switching, disability, culture, long relationships, and adversarial obfuscation. Severity is operational rather than clinical, and v0.1 lacks a formal information-needed state. Target validity measures taxonomy conformance, not accessible, consensual, locally available, or successful handoff. The incomplete informal editorial read is not validation; no response statistic, quotation, or relabeling is reported. No practitioner or youth validation has occurred. Child-centered is a design objective, not a measured property.")
    add_body(doc, "Before comparative or generalization claims, qualified practitioners must independently review the construct, action thresholds, target constraints, missing context, and jurisdictional feasibility. Youth advisors must separately review plausibility, agency, and accessibility without personal disclosure. Only then will independent authors create family-locked public and controlled-access tests without seeing the router.")
    add_body(doc, "Evidence gates version practitioner revisions and retain disagreement; keep youth and professional judgments separate; reserve a controlled-access family-locked test; and audit any response mapper against independent human coding. Failure at any gate blocks claims that depend on the next one, and no gate alone certifies deployment safety.", "Evidence gates. ")
    add_body(doc, "ChildEsc's practical use is narrow: an executable object for debating escalation thresholds, recipient defaults, and missing evidence. It cannot determine clinical risk, select a real person, assess response quality, or authorize disclosure.")
    add_body(doc, "ChildEsc supports audits without real child conversations, but could be misused as a safety certificate, embedded in coercive monitoring or disclosure, or encode culturally inappropriate routes that reduce youth agency. We prohibit automated contact or disclosure, retain disagreement, and gate operational use on practitioner, youth, and jurisdiction-specific review.", "Broader impacts. ")
    add_body(doc, "ChildEsc adds recipient-constrained routing as a machine-checkable object beyond action-only evaluation. The recurring gap shows the diagnostic value of asking when to escalate and who gets involved, without validating provisional routes. The workshop-facing question is how to represent missing context and practitioner/youth disagreement without turning routes into mandates. ChildEsc offers a technically testable starting point for the workshop's evaluation and ecosystem-safeguard agenda.", "Conclusion. ")

    add_heading(doc, "References")
    references = [
        "Arif, S., Borah, A., & Mihalcea, R. (2026). The Age of Curiosity Meets the Age of AI: Benchmarking Child Safety in Large Language Models. arXiv:2605.25510.",
        "Cha, H., Shukla, N., Barocas, S., Chouldechova, A., Kim, E., & Vaughan, J. W. (2026). Beyond 'I Can't Help With That': How Child Safety Experts Evaluate AI Chatbot Safety. arXiv:2608.07902. To appear at AIES 2026.",
        "Hua, Y., Na, H., Li, Z., Liu, F., Fang, X., Clifton, D., et al. (2025). A scoping review of large language models for generative tasks in mental health care. npj Digital Medicine, 8, 230.",
        "Hua, Y., Na, H., & Ayubcha, C. (2026). CARE-Bench: Benchmarking Patient-Facing LLM Triage. arXiv:2608.03731.",
        "Bean, A. M., Kearns, R. O., Romanou, A., Hafner, F. S., et al. (2025). Measuring What Matters: Construct Validity in Large Language Model Benchmarks. NeurIPS 2025.",
        "Jiao, J., Afroogh, S., Chen, K., Murali, A., Atkinson, D., & Dhurandhar, A. (2025). Safe-Child-LLM: A developmental benchmark for evaluating LLM safety in child-AI interactions. arXiv:2506.13510.",
        "Khoo, S., Chua, G., & Shong, R. (2025). MinorBench: A hand-built benchmark for content-based risks for children. arXiv:2503.10242.",
        "Krishna-Kumar, K., Lau, E., Robinson, V., Caldwell, J., Issaka, S., Wang, S., et al. (2026). CAREBench: A Child-Safety Risk Benchmark for Language Models. arXiv:2606.29685.",
        "KORA Benchmark. (2026). KORA: Open-Source Benchmark for AI Child Safety. https://github.com/korabench/benchmark.",
        "OpenAI. (2026, August 18). Introducing ChatGPT for Teens: Built for learning, backed by protections. https://openai.com/index/chatgpt-for-teens/.",
        "Pombal, J., D'Eon, M., Guerreiro, N. M., Martins, P. H., Farinhas, A., & Rei, R. (2025). MindEval: Benchmarking language models on multi-turn mental health support. arXiv:2511.18491.",
        "Ragone, G., Buono, P., Good, J., & Lanzilotti, R. (2026). Do Children Trust AI, and Should They? Proceedings of CHI '26. https://doi.org/10.1145/3772318.3790765.",
        "UNICEF. (2025). Guidance on AI and Children 3.0.",
        "UNICEF. (2026). When AI Becomes a Friend: Child Rights Risks, Harms, and Regulatory Responses to AI Chatbots and Companions.",
        "NeurIPS 2026 Workshop on Child Safety in AI. (2026). Call for Papers. https://childsafety-ai.github.io/.",
        "AI Child Safety Initiative. (2025). Preventing AI-Generated Child Sexual Abuse Material Necessitates New Approaches to AI Safety. https://aichildsafety.github.io/.",
    ]
    for reference in references:
        add_reference(doc, reference)

    add_heading(doc, "Appendix A. Artifact and schema")
    add_body(doc, "The research package contains standard-library Python source, 53 tests for the frozen scientific and manuscript path, 26 provider-path tests, 20 source families, a source register, family rationale map, controlled authoring protocol, frozen contract probes, deterministic generation, item-level predictions and assertions, aggregate metrics, family-bootstrap intervals, target-failure decomposition, ablations, prospective validation materials, and orthographic checks. Running make reproduce regenerates all 80 items and every deterministic reported value. All eight scored prompted runs were replayed internally from content-addressed caches before packaging; raw provider responses and caches are excluded, so reviewers can verify hashes and deterministic analyses but cannot replay provider responses from the distributed package.", size=9.5, after=4, line=10)

    add_heading(doc, "Appendix B. Metric details")
    add_body(doc, "Under-escalation is conditioned on items requiring human or urgent support. Strict valid-target rate requires a nonempty predicted set that is a subset of permitted targets and disjoint from forbidden targets. Exact route additionally requires the exact action and requires no emitted target for targetless actions. The action-route gap equals exact action accuracy minus exact-route accuracy. Action instability is the fraction of items whose predicted action changes across trials; RouteVar is the fraction whose action-target tuple changes. The interval CSV reports per-trial 95% family-bootstrap intervals from 1,000 resamples with seed 20260829.", size=9.5, after=4, line=10)

    add_heading(doc, "Appendix C. Diagnostic failure analysis")
    add_body(doc, "Agree is action agreement with the full-conversation run. Under is conditioned on escalation-required items; Valid on items with permitted targets; Unsafe on items with forbidden targets. These are policy-exposed diagnostics, not generalization estimates.", size=9.5, after=4, line=10)
    diagnostic_headers = ["Condition", "Agree", "Urgent", "Under", "Valid", "Unsafe", "Exact"]
    diagnostic_data = [
        ["Full conversation", "1.000", ".757", ".328", ".552", ".000", ".538"],
        ["Latest turn only", ".338", ".000", "1.000", ".000", ".000", ".138"],
        ["Adult inference off", "1.000", ".757", ".328", ".431", ".583", ".488"],
        ["Uppercase", "1.000", ".757", ".328", ".552", ".000", ".538"],
        ["Extra spaces", ".450", ".108", ".897", ".034", ".083", ".175"],
        ["Expanded contractions", ".975", ".730", ".362", ".534", ".000", ".525"],
    ]
    diagnostic_table = doc.add_table(rows=1, cols=len(diagnostic_headers))
    diagnostic_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    diagnostic_table.autofit = False
    mark_table_header(diagnostic_table.rows[0])
    diagnostic_widths = [Inches(1.45)] + [Inches(0.75)] * 6
    for index, header in enumerate(diagnostic_headers):
        cell = diagnostic_table.rows[0].cells[index]
        cell.width = diagnostic_widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "E7E7E7")
        set_cell_margins(cell, start=30, end=30)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cell.paragraphs[0].add_run(header), size=7.5, bold=True)
    for values in diagnostic_data:
        row = diagnostic_table.add_row()
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = diagnostic_widths[index]
            set_cell_margins(cell, start=30, end=30)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_font(cell.paragraphs[0].add_run(value), size=7.5)
    set_table_borders(diagnostic_table)

    add_heading(doc, "Appendix D. Prospective validation protocol")
    add_body(doc, "No authorized practitioner or youth research data collection occurred. After the technical evidence freeze, one non-professional individual with mental-health knowledge provided an incomplete informal editorial read. We treated it as manuscript-design feedback, not study data or Gate 1 evidence, and report no item-level response, quotation, agreement statistic, or label revision. Subject to an appropriate external determination, practitioners would independently label action, permitted and forbidden targets, missing context, and coercion risk. Youth advisors would separately review plausibility, agency, accessibility, and handoff framing without personal disclosure. Disagreement and pre-review labels would be retained. Protocol preparation is not validation.", size=9.5, after=4, line=10)

    add_heading(doc, "Appendix E. Independent holdout status")
    add_body(doc, "No independently authored holdout is reported. The prospective procedure requires an author who has not inspected router rules, existing scenario text, or item-level failures. Current authors and coding agents are ineligible. A future holdout must be evaluated once without router tuning and reported separately from v0.1.", size=9.5, after=4, line=10)

    add_heading(doc, "Appendix F. Compute, ethics, and LLM usage")
    add_body(doc, "Deterministic experiments use no GPU or training and complete in under one second on a laptop-class CPU. Reported-system smokes, complete trials, and the retained incomplete trial made 732 OpenRouter calls and cost $2.635; earlier failed development pilots cost another $0.215. Calls used no tools and generated only routing JSON, not counseling text. A general-purpose LLM coding assistant drafted candidate scenarios and initial labels, implementation, literature notes, and manuscript text. Prompted systems generated reported routing decisions, but no LLM judged outputs. Public release must retain content warnings and must not support automated data disclosure, emergency or caregiver contact, or claims of clinical validation.", size=9.5, after=4, line=10)

    add_heading(doc, "Appendix G. Source traceability and future authoring")
    add_body(doc, "The source register records each source's support and limits; v0.1 mappings are retrospective and domain-level, not item-level grounding. Future authoring requires controlled counterfactuals, timing and locale metadata, sourced rationales, non-procedural content, independent labels without router outputs, retained disagreement, and family-locked splits.", size=9.5, after=4, line=10)

    footer = section.footer.paragraphs[0]
    footer.clear()
    draft_run = footer.add_run("Anonymous collaboration draft | ")
    set_font(draft_run, size=8)
    add_page_number(footer)

    core = doc.core_properties
    core.title = "ChildEsc Workshop Paper"
    core.subject = "NeurIPS 2026 Child Safety in AI workshop submission draft"
    core.author = "Anonymous"
    core.keywords = "child safety, conversational AI, escalation, benchmark, human handoff"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
